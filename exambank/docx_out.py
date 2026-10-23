# -*- coding: utf-8 -*-
"""نمای خواندنی docx راست‌به‌چپ (اختیاری) از روی pages/*.json؛ قالب عنوان‌ها مثل نسخه‌ی قدیمی."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

PERSIAN_FONT = "B Nazanin"


def _rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))


def _font(run, size=12, bold=False):
    run.font.rtl = True
    run.font.size = Pt(size)
    run.font.name = PERSIAN_FONT
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    for a in ("w:cs", "w:ascii", "w:hAnsi"):
        rf.set(qn(a), PERSIAN_FONT)


def _line(doc, text, size=12, bold=False):
    p = doc.add_paragraph()
    _rtl(p)
    _font(p.add_run(text), size, bold)


def page_lines(rec):
    lines = []
    for q in rec["questions"]:
        mark = "(ادامه) " if q["continues_from_previous_page"] else ""
        lines.append(f"{q['number'] if q['number'] is not None else ''}) {mark}{q['stem']}".strip())
        for o in q["options"]:
            lines.append(f"   {o['label']}) {o['text']}")
        if q["continues_on_next_page"]:
            lines.append("   [ادامه در صفحه‌ی بعد]")
    for a in rec["answers"]:
        s = f"پاسخ {a['number']}: {'، '.join(a['correct']) or '—'}"
        if a["status"]:
            s += f" ({a['status']})"
        lines.append(s)
        if a["explanation"]:
            lines.append(f"   {a['explanation']}")
    for f in rec["figures"]:
        lines.append(f"[{f['kind']} {f['figure_id']}] {f['description_fa']}")
    return lines


def write_docx(path, title, source, page_records):
    doc = Document()
    _line(doc, f"سوالات استخراج‌شده — پروژه: {title}", 15, True)
    _line(doc, f"منبع: {source}", 10)
    for rec in sorted(page_records, key=lambda r: (r["pdf"], r["page"])):
        _line(doc, f"--- سوالات صفحه {rec['page']} ---", 13, True)
        for ln in page_lines(rec):
            _line(doc, ln)
    doc.save(path)
